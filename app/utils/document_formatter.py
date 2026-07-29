"""
Document formatting utilities for professional manuscript export.

Builds editable DOCX manuscripts with title page, table of contents, and
chapter formatting. Parses lightweight markdown-style markup in body text.
"""
from __future__ import annotations

import re
from typing import Iterable, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.schemas.export import BookChapter, StructuredBook

# ── Typography ────────────────────────────────────────────────────────────────
FONT_NAME = "Times New Roman"
BODY_SIZE = Pt(12)
CHAPTER_TITLE_SIZE = Pt(18)
SECTION_TITLE_SIZE = Pt(14)
TOC_TITLE_SIZE = Pt(16)
TITLE_PAGE_TITLE_SIZE = Pt(28)
TITLE_PAGE_SUBTITLE_SIZE = Pt(16)
TITLE_PAGE_AUTHOR_SIZE = Pt(14)
LINE_SPACING = 1.5

# Regex patterns for inline formatting
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*|__(.+?)__")
_ITALIC_RE = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)|_(.+?)_")
_BLOCKQUOTE_RE = re.compile(r"^>\s?(.*)$")
_HEADING_RE = re.compile(r"^#{1,6}\s+(.+)$")
_BULLET_RE = re.compile(r"^[-*•]\s+(.+)$")
_NUMBERED_RE = re.compile(r"^\d+[.)]\s+(.+)$")


def _set_run_font(run, size: Pt = BODY_SIZE, bold: bool = False, italic: bool = False) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = size
    run.bold = bold
    run.italic = italic


def _apply_paragraph_spacing(paragraph, before: int = 0, after: int = 6) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def _set_line_spacing(paragraph, multiplier: float = LINE_SPACING) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = multiplier


def _add_page_break(document: Document) -> None:
    document.add_page_break()


def _add_empty_lines(document: Document, count: int) -> None:
    for _ in range(count):
        p = document.add_paragraph("")
        _set_line_spacing(p)


def _add_formatted_runs(paragraph, text: str, base_size: Pt = BODY_SIZE) -> None:
    """Parse inline **bold** and *italic* markers into Word runs."""
    if not text:
        return

    pos = 0
    pattern = re.compile(
        r"(\*\*.+?\*\*|__.+?__|\*(?!\*).+?(?<!\*)\*(?!\*)|_.+?_)"
    )

    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            _set_run_font(run, size=base_size)

        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, size=base_size, bold=True)
        elif token.startswith("__") and token.endswith("__"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, size=base_size, bold=True)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, size=base_size, italic=True)
        elif token.startswith("_") and token.endswith("_"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, size=base_size, italic=True)
        pos = match.end()

    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_run_font(run, size=base_size)


def _add_body_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    _set_line_spacing(paragraph)
    _apply_paragraph_spacing(paragraph, before=0, after=8)
    _add_formatted_runs(paragraph, text.strip())


def _add_list_paragraph(
    document: Document,
    text: str,
    *,
    numbered: bool = False,
    level: int = 0,
) -> None:
    style = "List Number" if numbered else "List Bullet"
    try:
        paragraph = document.add_paragraph(style=style)
    except KeyError:
        paragraph = document.add_paragraph()
        prefix = f"{level + 1}. " if numbered else "• "
        text = prefix + text

    _set_line_spacing(paragraph)
    _apply_paragraph_spacing(paragraph, before=0, after=4)
    paragraph.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    _add_formatted_runs(paragraph, text.strip())


def _add_blockquote(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    _set_line_spacing(paragraph)
    _apply_paragraph_spacing(paragraph, before=4, after=8)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    run = paragraph.add_run(text.strip())
    _set_run_font(run, italic=True)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def _add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    _set_line_spacing(paragraph)
    _apply_paragraph_spacing(paragraph, before=12, after=6)
    run = paragraph.add_run(text.strip())
    _set_run_font(run, size=SECTION_TITLE_SIZE, bold=True)


def _render_content_block(document: Document, content: str) -> None:
    """Render multi-line content with lists, quotes, and inline formatting."""
    if not content or not content.strip():
        return

    lines = content.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        quote_match = _BLOCKQUOTE_RE.match(line)
        if quote_match:
            _add_blockquote(document, quote_match.group(1))
            i += 1
            continue

        heading_match = _HEADING_RE.match(line)
        if heading_match:
            _add_section_heading(document, heading_match.group(1))
            i += 1
            continue

        bullet_match = _BULLET_RE.match(line)
        if bullet_match:
            _add_list_paragraph(document, bullet_match.group(1), numbered=False)
            i += 1
            continue

        numbered_match = _NUMBERED_RE.match(line)
        if numbered_match:
            _add_list_paragraph(document, numbered_match.group(1), numbered=True)
            i += 1
            continue

        _add_body_paragraph(document, line)
        i += 1


def _configure_document_defaults(document: Document) -> None:
    """Apply normal margins and default body style."""
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = LINE_SPACING


def _add_title_page(document: Document, book: StructuredBook) -> None:
    _add_empty_lines(document, 8)

    title_p = document.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_line_spacing(title_p)
    title_run = title_p.add_run(book.title)
    _set_run_font(title_run, size=TITLE_PAGE_TITLE_SIZE, bold=True)

    if book.subtitle:
        subtitle_p = document.add_paragraph()
        subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_line_spacing(subtitle_p)
        _apply_paragraph_spacing(subtitle_p, before=12, after=0)
        subtitle_run = subtitle_p.add_run(book.subtitle)
        _set_run_font(subtitle_run, size=TITLE_PAGE_SUBTITLE_SIZE, italic=True)

    if book.author:
        author_p = document.add_paragraph()
        author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_line_spacing(author_p)
        _apply_paragraph_spacing(author_p, before=24, after=0)
        author_run = author_p.add_run(book.author)
        _set_run_font(author_run, size=TITLE_PAGE_AUTHOR_SIZE)

    _add_page_break(document)


def _add_table_of_contents(document: Document, chapters: Iterable[BookChapter]) -> None:
    toc_title = document.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_line_spacing(toc_title)
    _apply_paragraph_spacing(toc_title, before=0, after=18)
    toc_run = toc_title.add_run("Table of Contents")
    _set_run_font(toc_run, size=TOC_TITLE_SIZE, bold=True)

    for index, chapter in enumerate(chapters, start=1):
        entry = document.add_paragraph()
        _set_line_spacing(entry)
        _apply_paragraph_spacing(entry, before=0, after=6)
        entry.paragraph_format.left_indent = Inches(0.25)
        run = entry.add_run(f"Chapter {index}: {chapter.title}")
        _set_run_font(run, size=BODY_SIZE)

    _add_page_break(document)


def _add_chapter(
    document: Document,
    chapter_number: int,
    chapter: BookChapter,
) -> None:
    label_p = document.add_paragraph()
    _set_line_spacing(label_p)
    _apply_paragraph_spacing(label_p, before=0, after=6)
    label_run = label_p.add_run(f"Chapter {chapter_number}")
    _set_run_font(label_run, size=CHAPTER_TITLE_SIZE, bold=True)

    title_p = document.add_paragraph()
    _set_line_spacing(title_p)
    _apply_paragraph_spacing(title_p, before=0, after=12)
    title_run = title_p.add_run(chapter.title)
    _set_run_font(title_run, size=CHAPTER_TITLE_SIZE, bold=True)

    if chapter.content and chapter.content.strip():
        _render_content_block(document, chapter.content)

    for section in chapter.sections:
        if section.heading.strip():
            _add_section_heading(document, section.heading)
        if section.content.strip():
            _render_content_block(document, section.content)


def build_manuscript_docx(book: StructuredBook) -> Document:
    """
    Build a professionally formatted Word manuscript from structured book data.

    Extension points (future): images, tables, references, footnotes can be
    added via new helpers called from _render_content_block / _add_chapter.
    """
    document = Document()
    _configure_document_defaults(document)
    _add_title_page(document, book)
    _add_table_of_contents(document, book.chapters)

    for index, chapter in enumerate(book.chapters, start=1):
        if index > 1:
            _add_page_break(document)
        _add_chapter(document, index, chapter)

    return document


def parse_content_into_sections(content: str) -> Tuple[str, List[Tuple[str, str]]]:
    """
    Split plain/markdown chapter content into an intro and section list.

    Recognises markdown headings (## Section) and ALL-CAPS standalone lines.
    """
    if not content or not content.strip():
        return "", []

    lines = content.splitlines()
    intro_lines: List[str] = []
    sections: List[Tuple[str, str]] = []
    current_heading: Optional[str] = None
    current_lines: List[str] = []

    def flush_section() -> None:
        nonlocal current_heading, current_lines
        if current_heading:
            sections.append((current_heading, "\n".join(current_lines).strip()))
        current_heading = None
        current_lines = []

    for line in lines:
        heading_match = _HEADING_RE.match(line.strip())
        if heading_match:
            flush_section()
            current_heading = heading_match.group(1).strip()
            continue

        stripped = line.strip()
        if (
            stripped
            and stripped == stripped.upper()
            and len(stripped.split()) <= 8
            and not stripped.endswith(".")
            and not _BULLET_RE.match(stripped)
            and not _NUMBERED_RE.match(stripped)
        ):
            flush_section()
            current_heading = stripped.title()
            continue

        if current_heading:
            current_lines.append(line)
        else:
            intro_lines.append(line)

    flush_section()

    if not sections:
        return content.strip(), []

    return "\n".join(intro_lines).strip(), sections
